"""Append newly observed facility names to the manual dictionary, never decide them.

The same functions serve local source refreshes and the repository workflow.
Only explicit hospital/workplace columns are collected. Existing CSV bytes are
preserved, and each new row contains the literal source name and two blank cells.
"""
from __future__ import annotations

import argparse
import codecs
import csv
import hashlib
import importlib.util
import io
import json
import os
import re
import sys
import tempfile
import time
from pathlib import Path

CACHE_SCHEMA = 1
HEADER_OPTIONS = (
    {"raw name"}, {"final medical centre names", "final name"},
    {"standartized type medical centre", "standardized type medical centre", "final type"},
)
PLACEHOLDERS = {"", "-", "—", "–", "null", "none", "nan", "n/a", "not recorded", "няма"}


class ReferenceChangedError(RuntimeError):
    """A manual edit arrived while names were being collected or appended."""


def file_sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_reference_bytes(raw):
    # Preserve bytes separately: parsing is only for validation and deduplication.
    csv.field_size_limit(min(sys.maxsize, 2**31 - 1))
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"), newline="")))
    if not rows or len(rows[0]) != 3 or any(
        value.strip().casefold() not in allowed
        for value, allowed in zip(rows[0], HEADER_OPTIONS)
    ):
        raise ValueError("The dictionary must have Raw Name, Final Name, Final Type in that order "
                         "(the original longer header labels are also supported).")
    if any(len(row) != 3 for row in rows):
        raise ValueError("Every dictionary row must contain exactly three cells.")
    return rows[1:]


def append_names(reference, names, *, expected_sha256=None, dry_run=False):
    """Append missing spellings, preserving all existing bytes and manual decisions.

    Pass the hash captured before source collection as expected_sha256 to refuse
    overwriting a newer manual dictionary. Repeating a successful call is a no-op.
    No fuzzy normalization or automatic name/type decisions are performed.
    """
    reference = Path(reference).resolve()
    original = reference.read_bytes()
    before = hashlib.sha256(original).hexdigest()
    if expected_sha256 and before != expected_sha256:
        raise ReferenceChangedError("The facility dictionary changed; collect again against its latest version.")
    rows = _read_reference_bytes(original)
    known = {row[0].strip() for row in rows}
    observed = {name.strip() for name in names if isinstance(name, str)
                and name.strip().casefold() not in PLACEHOLDERS}
    added = sorted(observed - known)
    pending = sum(not row[1].strip() for row in rows)
    summary = {"reference": str(reference), "referenceSha256Before": before,
               "referenceSha256After": before, "existingRows": len(rows),
               "addedRows": len(added), "addedNames": added, "totalRows": len(rows) + len(added),
               "pendingRows": pending + len(added), "changed": bool(added), "dryRun": bool(dry_run),
               "newRowsHaveBlankNameAndType": True}
    if not added:
        return summary
    # Use the existing line separator without rewriting quoting, BOM or old rows.
    match = re.search(br"\r\n|\r|\n", original)
    separator = match.group().decode("ascii") if match else "\n"
    new_text = io.StringIO(newline="")
    csv.writer(new_text, lineterminator=separator).writerows((name, "", "") for name in added)
    boundary = b"" if original.endswith((b"\n", b"\r")) else separator.encode("ascii")
    proposed = original + boundary + new_text.getvalue().encode("utf-8")
    # Round-trip check includes embedded quotes and multiline literal source names.
    checked = _read_reference_bytes(proposed)
    if checked[:len(rows)] != rows or checked[len(rows):] != [[name, "", ""] for name in added]:
        raise ValueError("Dictionary append verification failed; the dictionary was not changed.")
    summary["referenceSha256After"] = hashlib.sha256(proposed).hexdigest()
    if dry_run:
        return summary
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(prefix="." + reference.name + ".", suffix=".pending",
                                         dir=reference.parent, delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(proposed)
            stream.flush()
            os.fsync(stream.fileno())
        # A file changed while the source scan ran must be retried, not replaced.
        for attempt, delay in enumerate((0.05, 0.1, 0.25, 0.5, 1.0, 0.0)):
            if file_sha256(reference) != before:
                raise ReferenceChangedError("The facility dictionary changed during append; retry against the current file.")
            try:
                temporary.replace(reference)
                break
            except OSError as exc:
                # OneDrive/Excel/antivirus may hold a short-lived file handle.
                # Retry only sharing/permission locks; all other errors surface.
                if attempt == 5 or not (isinstance(exc, PermissionError) or getattr(exc, "winerror", None) in {5, 32, 33}):
                    raise
                time.sleep(delay)
        temporary = None
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return summary


def _inside(root, value, *, allow_external=False):
    path = Path(value)
    resolved = path.resolve() if path.is_absolute() else (root / path).resolve()
    if not allow_external and not resolved.is_relative_to(root):
        raise ValueError(f"Manifest path escapes the repository: {value!r}")
    return resolved


def _collector(root, manifest, *, allow_external=False):
    candidates = []
    if manifest.get("collector"):
        candidates.append(_inside(root, manifest["collector"], allow_external=allow_external))
    candidates.extend((Path(__file__).with_name("collect_facility_dictionary_names.py"),
                       root / "tools/maintenance/collect_facility_dictionary_names.py"))
    path = next((path for path in candidates if path.is_file()), None)
    if path is None:
        raise FileNotFoundError("Package collect_facility_dictionary_names.py beside this script "
                                "or configure its repository-relative collector path.")
    spec = importlib.util.spec_from_file_location("facility_dictionary_name_collector", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module, file_sha256(path)


def _cache_names(path, cache_key):
    try:
        # Names-only caches are small; reject corrupt/unbounded files before loading.
        if path.stat().st_size > 32 * 1024 * 1024:
            return None
        payload = json.loads(path.read_text(encoding="utf-8"))
        names = payload.get("names")
        if (payload.get("key") == cache_key and isinstance(names, list)
                and all(isinstance(name, str) for name in names)):
            return names
    except (OSError, ValueError, TypeError):
        pass
    return None


def _save_cache(path, cache_key, names):
    path.parent.mkdir(parents=True, exist_ok=True)
    content = json.dumps({"key": cache_key, "names": sorted(names)}, ensure_ascii=False,
                         separators=(",", ":")).encode("utf-8")
    if len(content) > 32 * 1024 * 1024:
        return
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(prefix=".names-", suffix=".pending", dir=path.parent,
                                         delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(content)
        temporary.replace(path)
        temporary = None
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _csv_encoding(path, configured):
    if configured != "auto":
        return configured
    # Incremental decoding does not mistake a split multibyte character for cp1251.
    decoder = codecs.getincrementaldecoder("utf-8-sig")()
    try:
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                decoder.decode(chunk, final=False)
        decoder.decode(b"", final=True)
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1251"


def _source_inventory(root, manifest):
    entries = list(manifest["inputs"])
    known = {str(Path(entry["path"]).as_posix()) for entry in entries}
    if manifest.get("discover_outputs"):
        # Future scraper outputs are read only through explicit workplace columns.
        # A future plain Name/Име field is NOT assumed to mean a hospital.
        for path in sorted(root.glob("**/*_outputs/**/*")):
            if not path.is_file() or path.suffix.lower() not in {".csv", ".xlsx", ".xlsm"}:
                continue
            relative = path.relative_to(root).as_posix()
            if relative in known:
                continue
            known.add(relative)
            entries.append({"path": relative, "source": path.parent.name,
                            "required": False, "discovered": True})
    return entries


def collect_manifest_names(root, manifest, *, cache_dir=None, allow_external=False):
    """Read explicitly inventoried sources with safe optional names-only caching.

    Each input specifies path, source, format, and optional sheets/facility_sheets.
    A generic Name column is used only with facility_table=true or an explicit
    facility sheet, so doctor names and descriptions cannot become hospital names.
    Missing required inputs fail visibly; optional not-yet-published inputs report
    their absence. A changed file is never represented by a stale cached list.
    """
    root = Path(root).resolve()
    csv.field_size_limit(min(sys.maxsize, 2**31 - 1))
    if isinstance(manifest, (str, Path)):
        manifest = json.loads(Path(manifest).read_text(encoding="utf-8-sig"))
    if manifest.get("schema") != 1 or not isinstance(manifest.get("inputs"), list):
        raise ValueError("Unsupported facility dictionary input manifest.")
    collector_module, collector_hash = _collector(root, manifest, allow_external=allow_external)
    all_names, reports = set(), []
    cache_dir = Path(cache_dir).resolve() if cache_dir else None
    for entry in _source_inventory(root, manifest):
        if not isinstance(entry, dict) or not entry.get("path"):
            raise ValueError("Every manifest input needs a path.")
        path = _inside(root, entry["path"], allow_external=allow_external)
        if not path.is_file():
            if entry.get("required", True):
                raise FileNotFoundError(f"Required facility-name source is missing: {entry['path']}")
            reports.append({"path": entry["path"], "status": "optional_missing", "names": 0})
            continue
        kind = entry.get("format") or path.suffix.lstrip(".").lower()
        encoding = _csv_encoding(path, entry.get("encoding", "utf-8-sig")) if kind == "csv" else None
        delimiter = entry.get("delimiter", ",")
        if kind == "csv" and entry.get("discovered"):
            # Business-registry outputs can be enormous and contain no hospitals.
            # Inspect their header before hashing or iterating irrelevant rows.
            with path.open(encoding=encoding, newline="") as stream:
                headers = next(csv.reader(stream, delimiter=delimiter), [])
            if not any(collector_module.is_name_column(column) or column in collector_module.JSON_COLUMNS
                       for column in headers):
                reports.append({"path": entry["path"], "status": "no_facility_columns", "names": 0})
                continue
        digest = file_sha256(path)
        key_data = {"schema": CACHE_SCHEMA, "collector": collector_hash, "file": digest, "input": entry}
        cache_key = hashlib.sha256(json.dumps(key_data, sort_keys=True, ensure_ascii=False).encode("utf-8")).hexdigest()
        cache_path = cache_dir / (hashlib.sha256(str(path).encode("utf-8")).hexdigest() + ".json") if cache_dir else None
        names = _cache_names(cache_path, cache_key) if cache_path else None
        hit = names is not None
        warnings, scanned_rows = [], 0
        if not hit:
            collector = collector_module.NameCollector()
            if kind == "csv":
                try:
                    origin = str(path.relative_to(root))
                except ValueError:
                    origin = str(path)
                with path.open(encoding=encoding, newline="") as stream:
                    for row in csv.DictReader(stream, delimiter=delimiter):
                        collector.table_row(row, origin, source=entry.get("source", ""),
                                            facility_table=entry.get("facility_table", False))
                        collector.scanned[origin] += 1
            elif kind in {"xlsx", "xlsm"}:
                facility_sheets = entry.get("facility_sheets", ())
                if entry.get("facility_table"):
                    facility_sheets = entry.get("sheets")
                    if not facility_sheets:
                        from openpyxl import load_workbook
                        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
                        try:
                            facility_sheets = workbook.sheetnames
                        finally:
                            workbook.close()
                collector.workbook_file(path, root, source=entry.get("source", ""),
                                        sheets=entry.get("sheets"),
                                        facility_sheets=facility_sheets)
            else:
                raise ValueError(f"Unsupported manifest source format: {kind!r}")
            # Reject a source modified during collection instead of caching partial evidence.
            if file_sha256(path) != digest:
                raise ReferenceChangedError(f"Source changed during collection: {entry['path']}")
            names = sorted(collector.names)
            warnings, scanned_rows = sorted(set(collector.warnings)), sum(collector.scanned.values())
            # A required Excel sheet disappearing must not silently reduce coverage.
            if any("configured worksheet" in warning for warning in warnings) and entry.get("required", True):
                raise ValueError(f"Facility-name source needs attention: {entry['path']}: {'; '.join(warnings)}")
            if cache_path and not warnings:
                _save_cache(cache_path, cache_key, names)
        if file_sha256(path) != digest:
            raise ReferenceChangedError(f"Source changed during collection: {entry['path']}")
        all_names.update(names)
        reports.append({"path": entry["path"], "status": "cached" if hit else "scanned",
                        "sha256": digest, "names": len(names), "scannedRows": scanned_rows, "warnings": warnings})
    return {"names": sorted(all_names), "cacheHits": sum(row["status"] == "cached" for row in reports),
            "scannedFiles": sum(row["status"] == "scanned" for row in reports), "inputs": reports}


def sync_manifest(root, manifest_path, *, reference=None, cache_dir=None, dry_run=False):
    root = Path(root).resolve()
    manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8-sig"))
    reference = _inside(root, reference or manifest["reference"])
    expected = file_sha256(reference)
    observed = collect_manifest_names(root, manifest, cache_dir=cache_dir)
    summary = append_names(reference, observed["names"], expected_sha256=expected, dry_run=dry_run)
    return {**summary, "observedLiteralNames": len(observed["names"]),
            "cacheHits": observed["cacheHits"], "scannedFiles": observed["scannedFiles"],
            "inputs": observed["inputs"]}


def manifest_from_config(config, *, repository=None):
    """Freeze enabled GitHub copy rules into the cloud's explicit input inventory.

    This export contains repository paths only, never local personal paths or
    credentials. Regenerate it when adding/removing a GitHub source in config.
    """
    manifests = []
    for repository_id, repo in config.get("github_source_repos", {}).items():
        if repo.get("enabled") is False or repository and repository_id != repository:
            continue
        inputs, reference = [], None
        for rule in repo.get("copy_rules", []):
            if rule.get("enabled") is False:
                continue
            config_id = rule.get("config_path_id", "")
            path = rule.get("source", "")
            if config_id == "facility_name_reference.path":
                reference = path
                continue
            if Path(path).suffix.lower() not in {".csv", ".xlsx", ".xlsm"}:
                continue
            entry = {"path": path, "source": rule.get("name", repository_id), "required": True}
            spec = config
            for part in config_id.split(".")[:-1] if config_id else []:
                spec = spec.get(part, {}) if isinstance(spec, dict) else {}
            if config_id.endswith(".hospitals_path"):
                entry["facility_table"] = True
            if isinstance(spec, dict):
                if spec.get("sheet"):
                    entry["sheets"] = [spec["sheet"]]
                selected = [spec[field] for field in ("bls_sheet", "managers_sheet", "hospitals_sheet", "addresses_sheet", "doctors_sheet") if spec.get(field)]
                if selected:
                    entry["sheets"] = list(dict.fromkeys(selected))
                facilities = [spec[field] for field in ("managers_sheet", "hospitals_sheet") if spec.get(field)]
                if facilities:
                    entry["facility_sheets"] = facilities
            inputs.append(entry)
        if reference:
            manifests.append({"schema": 1, "reference": reference, "discover_outputs": True,
                              "collector": ".github/scripts/facility-dictionary/collect_facility_dictionary_names.py",
                              "inputs": inputs})
    if len(manifests) != 1:
        raise ValueError("Choose exactly one configured GitHub repository with a facility dictionary.")
    return manifests[0]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path.cwd())
    parser.add_argument("--manifest", type=Path, default=Path(__file__).with_name("sources.json"))
    parser.add_argument("--reference")
    parser.add_argument("--cache-dir", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--report", type=Path)
    parser.add_argument("--export-config", type=Path, help="Export a cloud manifest from config/sources.json, then exit.")
    parser.add_argument("--repository", help="Configured GitHub repository ID for --export-config.")
    args = parser.parse_args()
    if args.export_config:
        payload = manifest_from_config(json.loads(args.export_config.read_text(encoding="utf-8-sig")), repository=args.repository)
        args.manifest.parent.mkdir(parents=True, exist_ok=True)
        args.manifest.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({"manifest": str(args.manifest), "inputs": len(payload["inputs"])}, indent=2))
        return
    summary = sync_manifest(args.root, args.manifest, reference=args.reference,
                            cache_dir=args.cache_dir, dry_run=args.dry_run)
    report = {key: value for key, value in summary.items() if key != "addedNames"}
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
