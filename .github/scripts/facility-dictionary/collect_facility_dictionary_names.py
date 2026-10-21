"""Preview a complete, editable facility-name dictionary without changing inputs.

The upstream three-column dictionary remains authoritative. Existing rows,
including duplicates and empty cells, are copied verbatim at cell level. Only
new literal spellings are appended, with empty standardized name/type cells.
Descriptions are never mined for guessed hospital names or employment links.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from functools import lru_cache
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
HEADERS = ('Raw Name', 'Final Medical Centre Names', 'Standartized Type Medical Centre')
HEADER_OPTIONS = (
    {'raw name'}, {'final medical centre names', 'final name'},
    {'standartized type medical centre', 'standardized type medical centre', 'final type'},
)
MAP_NAME_FIELDS = ('name', 'sourceName', 'originalName', 'standardName',
                   'standardizedName', 'manualName')
TABLE_NAMES = (
    'doctors_master_final', 'doctors_master_auto', 'candidate_match_audit',
    'manual_review_candidates', 'automatic_candidate_matches',
    'trusted_automatic_candidate_matches', 'base_linked_manual_candidates',
    'dropped_candidates', 'his_doctor_hospital_details', 'his_hospital_doctor_details',
    'his_hospital_workplace_summary', 'his_doctor_hospital_summary',
    'his_doctor_workplace_view', 'credoweb_facilities', 'credoweb_source_evidence',
    'bestdoctors_source_evidence', 'doctor_address_locations', 'doctor_manual_workplaces',
    'hospital_map_locations', 'source_profile_locations', 'facility_name_reference_audit', 'canonical_facilities',
    'canonical_facility_members',
)
NAME_COLUMNS = {
    'Doctor_Workplace_All', 'Doctor_Workplace_All_Source', 'Workplaces_All',
    'Workplaces_All_Source', 'Candidate_Workplace', 'HIS_Hospital_Name',
    'HIS_Final_Hospital_Names', 'HIS_Final_Hospital_Names_Source',
    'HIS_Usable_Hospital_Names', 'HIS_Review_Hospital_Names', 'Hospital_Name',
    'Original_Hospital_Name', 'Map_Hospital_Name', 'Standardized_Hospital_Name',
    'Facility_Name', 'Original_Name', 'Standardized_Name', 'Workplace',
    'Workplace_Name', 'Hospital', 'Recorded_Workplace', 'Facility_Previous_Names',
    'Doctor_Facility_Migration_Names',
    'recordedWorkplaces', 'Manual_Standard_Hospital_Name', 'Reference_Standard_Hospital_Name',
    'HIS_Hospital_Name_Source', 'HIS_Hospital_Names', 'HIS_Hospital_Names_Source',
    'HIS_Hospital_Names_Raw', 'HIS_Hospital_Standardized_Name', 'HIS_Hospital_Manual_Standard_Name',
    'HIS_Official_Hospital_Names', 'HIS_Manager_Potential_Hospital_Names', 'HIS_Review_Only_Hospital_Names',
    'CredoWebFacility.Name',
}
JSON_COLUMNS = {
    'Candidate_CredoWeb_Workplaces_JSON', 'CredoWeb.Workplaces_JSON', 'Workplaces_JSON',
    'Doctor_Manual_Workplace_Locations', 'Source_Link_Data', 'Facility_Merge_Source_Values',
    'Doctor_Manual_Source_Values', 'Doctor_Source_Link_Baseline', 'Previous_Names_JSON',
    'Historical_Hospital_Names',
}
PLACEHOLDERS = {'', '-', '—', '–', 'null', 'none', 'nan', 'n/a', 'not recorded', 'няма'}
CONTEXT_FIELDS = ('cities', 'addresses', 'registryIds', 'recordIds', 'doctorUins',
                  'sources', 'sourceUrls', 'origins', 'kinds', 'nameFields')


def parts(value):
    """The pipeline's list delimiter is a spaced pipe, not arbitrary commas."""
    if isinstance(value, (list, tuple, set)):
        return [str(item).strip() for item in value if str(item).strip()]
    return [item.strip() for item in str(value or '').split(' | ') if item.strip()]


@lru_cache(maxsize=4096)
def is_name_column(column):
    if column in NAME_COLUMNS:
        return True
    key = column.casefold()
    return bool(re.fullmatch(r'.*\.(?:workplace(?:_unspecified|_historical)?|hospital(?:_\d+|_corrected|_short)?)', key)
                or re.fullmatch(r'hospital_\d+', key)
                or re.fullmatch(r'болница(?:_\d+)?', key)
                or key in {'болница corrected', 'месторабота', 'лечебно заведение', 'наименование на лечебното заведение',
                    'наименование на лз', 'hospital name', 'hospital_name', 'hospital_names',
                    'workplace name', 'workplace_name', 'institution_name'})


def _resolve(base, value):
    path = Path(value)
    return path if path.is_absolute() else (base / path).resolve()


def read_reference(path):
    """Do not normalize cells, collapse duplicate rows, or fill empty decisions."""
    with Path(path).open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.reader(stream)
        headers = next(reader, [])
        if len(headers) != 3 or any(value.strip().casefold() not in allowed
                                  for value, allowed in zip(headers, HEADER_OPTIONS)):
            raise ValueError('The facility reference must contain Raw Name, Final Name and Final Type '
                             '(or their original longer header labels), in that order.')
        rows = list(reader)
    if any(len(row) != 3 for row in rows):
        raise ValueError('Every facility reference row must contain exactly three cells.')
    return headers, rows


class NameCollector:
    """One context record per literal spelling; labels never establish identity."""
    def __init__(self):
        self.names = {}
        self.scanned = Counter()
        self.warnings = []

    def add(self, value, context=None, *, field=''):
        if not isinstance(value, str):
            return
        name = value.strip()
        if name.casefold() in PLACEHOLDERS:
            return
        item = self.names.setdefault(name, {key: set() for key in CONTEXT_FIELDS})
        for key, values in (context or {}).items():
            if key in item:
                item[key].update(parts(values))
        if field:
            item['nameFields'].add(field)
        if re.match(r'^(?:Workplace of |Месторабота на |Лекарски адрес)', name, re.I):
            item['kinds'].add('generated_doctor_workplace_label')

    def facility(self, record, origin, *, kind='map_facility'):
        if not isinstance(record, dict):
            return
        workplace = bool(record.get('isDoctorWorkplace') or
                         str(record.get('id') or '').startswith('doctor-location:'))
        context = {'cities': record.get('city', ''), 'addresses': record.get('addresses', []),
            'registryIds': parts(record.get('registryId')) + parts(record.get('registryIds'))
                + parts(record.get('sourceRegistryIds')) + parts(record.get('targetRegistryIds')),
            'recordIds': record.get('id', ''),
            'doctorUins': [str(row.get('uin') or '') for row in record.get('doctors', [])
                           if isinstance(row, dict)] + parts(record.get('ownerUin')),
            'sourceUrls': parts(record.get('sourceUrls')) + parts(record.get('sourceUrl')) + parts(record.get('urls')),
            'sources': record.get('sources', []), 'origins': origin,
            'kinds': ['doctor_workplace' if workplace else kind]}
        for field in MAP_NAME_FIELDS:
            self.add(record.get(field), context, field=field)
        for field in ('nameAliases', 'previousNames'):
            for name in parts(record.get(field)):
                self.add(name, context, field=field)
        for child in record.get('groupedFacilities', []):
            if isinstance(child, dict):
                self.facility(child, origin + '/groupedFacilities', kind='grouped_facility')
        for child in record.get('mergedSources', []):
            if isinstance(child, dict):
                self.facility(child, origin + '/mergedSources', kind='archived_facility')

    def table_row(self, row, origin, *, source='', facility_table=False):
        context = {'origins': origin, 'sources': parts(source) + parts(row.get('Source_Name')),
            'cities': [], 'addresses': [], 'registryIds': [], 'recordIds': [],
            'doctorUins': [], 'sourceUrls': [], 'kinds': ['source_facility' if facility_table else 'recorded_workplace']}
        groups = {
            'cities': ('City', 'Град', 'Town', 'Address_Cities', 'HIS_Hospital_City', 'Facility_City'),
            'addresses': ('Address', 'Адрес', 'Адреси', 'Candidate_Address', 'Candidate_Address_Raw',
                          'Doctor_Address_All', 'Doctor_Address_All_Raw', 'HIS_Hospital_Address', 'Full_Address',
                          'Workplace_Address'),
            'registryIds': ('Registry_ID', 'Registry_IDs', 'Hospital_ID', 'HIS_Hospital_ID', 'HIS_Hospital_IDs',
                            'Facility_Registry_IDs', 'Candidate_Hospital_IDs'),
            'recordIds': ('Hospital_Record_Key', 'Facility_Record_Key', 'Record_Key', 'Review_ID',
                          'Candidate_ID', 'Source_Profile_ID', 'Source_Key'),
            'doctorUins': ('Doctor_UIN', 'UIN', 'Candidate_UIN', 'Correct_BLS_UIN'),
            'sourceUrls': ('Source_URL', 'Source_URLs', 'Candidate_URL', 'URL', 'Източник', 'Name_Reference_Source'),
        }
        for key, columns in groups.items():
            for column in columns:
                context[key].extend(parts(row.get(column)))
        for column, value in row.items():
            if not isinstance(column, str):
                continue
            if is_name_column(column) or (facility_table and column in {'Name', 'Име', 'Наименование'}):
                values = value
                if column == 'Facility_Previous_Names' and str(value or '').lstrip().startswith('['):
                    try:
                        values = json.loads(value)
                    except (ValueError, TypeError):
                        pass
                for name in parts(values):
                    self.add(name, context, field=column)
            if column in JSON_COLUMNS and value:
                try:
                    nested = json.loads(value) if isinstance(value, str) else value
                except (ValueError, TypeError):
                    self.warnings.append(f'{origin}: invalid structured JSON in {column}')
                    continue
                if column in {'Historical_Hospital_Names', 'Previous_Names_JSON'} and isinstance(nested, list):
                    for name in nested:
                        self.add(name, context, field=column)
                else:
                    self.structured(nested, origin + '/' + column, context,
                        place_context=column in {'Candidate_CredoWeb_Workplaces_JSON', 'CredoWeb.Workplaces_JSON',
                            'Workplaces_JSON', 'Doctor_Manual_Workplace_Locations'})

    def structured(self, value, origin, context=None, *, place_context=False):
        """Read structured names/snapshots, never free-text description strings."""
        if isinstance(value, list):
            for child in value:
                self.structured(child, origin, context, place_context=place_context)
        elif isinstance(value, dict):
            ident = str(value.get('id') or '')
            facility_record = ('registryId' in value or value.get('isDoctorWorkplace') is True
                or ident.startswith(('doctor-location:', 'manual-facility:'))
                or isinstance(value.get('doctors'), list) and ('addresses' in value or 'city' in value))
            if facility_record:
                self.facility(value, origin, kind='saved_facility_snapshot')
            local = {**(context or {}), 'origins': origin}
            if place_context or facility_record:
                local = {**local,
                    'cities': parts(local.get('cities')) + parts(value.get('city')),
                    'addresses': parts(local.get('addresses')) + parts(value.get('address')) + parts(value.get('fullAddress')),
                    'recordIds': parts(local.get('recordIds')) + parts(value.get('id')),
                    'sourceUrls': parts(local.get('sourceUrls')) + parts(value.get('sourceUrl')),
                    'kinds': parts(local.get('kinds')) + ['structured_workplace']}
            # Doctor overrides also contain a plain `name` field. Read that
            # field only in an explicitly structured workplace/snapshot slot.
            if place_context or facility_record:
                self.add(value.get('name'), local, field='name')
            if place_context:
                self.add(value.get('label'), local, field='label')
            for field in ('workplace', 'workplaceName', 'hospitalName', 'Workplace'):
                self.add(value.get(field), local, field=field)
            merge_decision = 'targetId' in value and ('sourceRecord' in value or 'targetRecord' in value)
            if merge_decision:
                for field in ('sourceName', 'targetName', 'sourceNameBefore', 'targetNameBefore',
                              'targetManualNameBefore', 'finalName'):
                    self.add(value.get(field), local, field=field)
                for field in ('sourceNameAliases', 'targetNameAliases'):
                    for name in parts(value.get(field)):
                        self.add(name, local, field=field)
            for field, child in value.items():
                if is_name_column(field):
                    for name in parts(child):
                        self.add(name, local, field=field)
                elif field in JSON_COLUMNS and isinstance(child, str) and child:
                    try:
                        nested = json.loads(child)
                    except (ValueError, TypeError):
                        self.warnings.append(f'{origin}: invalid structured JSON in {field}')
                        continue
                    self.structured(nested, origin + '/' + field, local,
                        place_context=field in {'Candidate_CredoWeb_Workplaces_JSON', 'CredoWeb.Workplaces_JSON',
                            'Workplaces_JSON', 'Doctor_Manual_Workplace_Locations'})
                elif isinstance(child, (dict, list)):
                    child_context = local
                    if re.fullmatch(r'[0-9]{9,10}', field):
                        child_context = {**local, 'doctorUins': parts(local.get('doctorUins')) + [field.zfill(10)]}
                    self.structured(child, origin + '/' + field, child_context,
                        place_context=field in {'manualLocations', 'workplaces', 'groupedFacilities',
                            'sourceRecord', 'targetRecord', 'promotedOriginalWorkplaceRecord'}
                            or field == 'locations' and place_context)

    def csv_file(self, path, root, *, source='', facility_table=False, encoding='utf-8-sig'):
        if not path.exists():
            return
        try:
            origin = str(path.relative_to(root))
        except ValueError:
            origin = str(path)
        with path.open(encoding=encoding, newline='') as stream:
            reader = csv.DictReader(stream)
            for row in reader:
                self.table_row(row, origin, source=source, facility_table=facility_table)
                self.scanned[origin] += 1

    def workbook_file(self, path, root, *, source='', sheets=None, facility_sheets=()):
        """Read explicit workplace columns from configured raw Excel sources.

        A plain Name column is accepted only in a known facility sheet, never
        in a doctor worksheet. Looking for the header permits ordinary named
        tables whose header is below a title row without importing formulas.
        """
        if not path.exists():
            return
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError('openpyxl is required to cover configured raw Excel sources.') from exc
        try:
            origin_base = str(path.relative_to(root))
        except ValueError:
            origin_base = str(path)
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            wanted = list(sheets) if sheets else workbook.sheetnames
            for sheet in wanted:
                if sheet not in workbook.sheetnames:
                    self.warnings.append(f'{origin_base}: configured worksheet {sheet!r} is missing')
                    continue
                facility_table = sheet in facility_sheets
                origin = origin_base + '#' + sheet
                rows = workbook[sheet].iter_rows(values_only=True)
                header = None
                for row_number, cells in enumerate(rows, 1):
                    if header is None:
                        candidate = [str(value or '').strip() for value in cells]
                        named = any(is_name_column(column) for column in candidate)
                        known_facility = facility_table and any(column in {'Name', 'Име', 'Наименование'} for column in candidate)
                        if named or known_facility:
                            header = candidate
                        elif row_number >= 50:
                            break
                        continue
                    row = {column: '' if value is None else str(value)
                           for column, value in zip(header, cells) if column}
                    self.table_row(row, origin, source=source, facility_table=facility_table)
                    self.scanned[origin] += 1
        finally:
            workbook.close()


def collect(root=ROOT, reference=None):
    """Return the proposed dictionary and context. Read-only until JSON export."""
    root = Path(root).resolve()
    csv.field_size_limit(min(sys.maxsize, 2**31 - 1))
    config_path = root / 'config/sources.json'
    config = json.loads(config_path.read_text(encoding='utf-8-sig')) if config_path.exists() else {}
    if reference is None:
        spec = config.get('facility_name_reference') or {}
        if not spec.get('path'):
            raise ValueError('Configure facility_name_reference.path or pass --reference.')
        reference = _resolve(config_path.parent, spec['path'])
    else:
        reference = _resolve(root, reference)
    reference_digest = hashlib.sha256(reference.read_bytes()).hexdigest()
    headers, existing = read_reference(reference)
    collector = NameCollector()
    map_path = root / 'outputs/hospital_map/data.js'
    if map_path.exists():
        text = map_path.read_text(encoding='utf-8')
        data = json.loads(text.split('=', 1)[1].rstrip(';\r\n '))
        for hospital in data.get('hospitals', []):
            collector.facility(hospital, 'outputs/hospital_map/data.js')
        for hospital in data.get('facilityMergeRecords', {}).values():
            collector.facility(hospital, 'outputs/hospital_map/data.js/facilityMergeRecords', kind='archived_facility')
    for name in TABLE_NAMES:
        collector.csv_file(root / 'outputs/final' / (name + '.csv'), root,
            facility_table=name in {'credoweb_facilities', 'canonical_facilities', 'canonical_facility_members'})
    collector.csv_file(root / 'inputs/hospital_map_facility_overrides.csv', root)
    collector.csv_file(root / 'inputs/manual_link_corrections.csv', root)
    manual_workplaces = _resolve(config_path.parent, config['manual_workplace_addresses_csv']) \
        if config.get('manual_workplace_addresses_csv') else root / 'inputs/manual_workplace_addresses.csv'
    if manual_workplaces.exists():
        try:
            manual_workplaces.read_bytes().decode('utf-8-sig')
            encoding = 'utf-8-sig'
        except UnicodeDecodeError:
            encoding = 'cp1251'
        collector.csv_file(manual_workplaces, root, source='Manual workplace addresses', encoding=encoding)
    seen_paths = set()
    for source, spec in config.get('sources', {}).items():
        if not isinstance(spec, dict) or spec.get('enabled') is False:
            continue
        for field in ('path', 'hospitals_path', 'workplaces_path'):
            if not spec.get(field):
                continue
            path = _resolve(config_path.parent, spec[field])
            if path in seen_paths:
                continue
            seen_paths.add(path)
            if path.suffix.casefold() == '.csv':
                collector.csv_file(path, root, source=source, facility_table=field == 'hospitals_path')
            elif path.suffix.casefold() in {'.xlsx', '.xlsm'}:
                collector.workbook_file(path, root, source=source,
                    sheets=[spec['sheet']] if spec.get('sheet') else None)
    his_spec = config.get('his_doctor_hospitals') or {}
    if his_spec.get('enabled') is not False:
        for source, spec, sheet_fields in (
                ('HIS OpenData', his_spec, ('bls_sheet', 'managers_sheet')),
                ('HISBG', his_spec.get('hisbg_scraper') or {}, ('hospitals_sheet', 'addresses_sheet', 'doctors_sheet'))):
            if spec.get('enabled') is False or not spec.get('path'):
                continue
            path = _resolve(config_path.parent, spec['path'])
            sheets = list(dict.fromkeys(spec[field] for field in sheet_fields if spec.get(field)))
            if path.suffix.casefold() in {'.xlsx', '.xlsm'}:
                collector.workbook_file(path, root, source=source, sheets=sheets or None,
                    facility_sheets=[spec[field] for field in ('hospitals_sheet', 'managers_sheet') if spec.get(field)])
    for name in ('hospital_map_facility_merges.json', 'hospital_map_doctor_overrides.json'):
        path = root / 'inputs' / name
        if path.exists():
            collector.structured(json.loads(path.read_text(encoding='utf-8-sig')), str(path.relative_to(root)))
    # Every original spelling remains independent; case, punctuation, and inner
    # spaces may carry a useful correction decision and are never collapsed.
    known = {row[0].strip() for row in existing}
    added = [[name, '', ''] for name in sorted(collector.names) if name not in known]
    contexts = []
    for name, evidence in sorted(collector.names.items()):
        evidence['doctorUins'] = {value.zfill(10) if re.fullmatch(r'[0-9]{9,10}', value) else value
                                 for value in evidence['doctorUins']}
        contexts.append({'rawName': name, 'alreadyInReference': name in known,
            **{key: sorted(values) for key, values in evidence.items()}})
    context_headers = ['rawName', 'alreadyInReference', *CONTEXT_FIELDS]
    if hashlib.sha256(reference.read_bytes()).hexdigest() != reference_digest:
        raise ValueError('The reference changed during collection. Repeat against its current rows before publishing.')
    return {'headers': headers, 'existingRows': existing, 'addedRows': added,
        'allRows': existing + added, 'contextRows': contexts,
        'contextHeaders': context_headers,
        'summary': {'reference': str(reference), 'referenceSha256': reference_digest,
            'existingRows': len(existing),
            'existingUniqueLiteralNames': len(known), 'observedLiteralNames': len(collector.names),
            'addedRows': len(added), 'totalRows': len(existing) + len(added),
            'pendingExistingRows': sum(not row[1].strip() for row in existing),
            'newRowsHaveBlankNameAndType': True, 'scannedTables': dict(collector.scanned),
            'warnings': sorted(set(collector.warnings))}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=ROOT)
    parser.add_argument('--reference', type=Path)
    parser.add_argument('--output-json', type=Path,
        help='Write a review payload only; the live dictionary and databases are never modified.')
    parser.add_argument('--authoring-json', type=Path,
        help='Write the same three-column rows and summary without the larger per-name context.')
    args = parser.parse_args()
    result = collect(args.root, args.reference)
    if args.output_json and args.authoring_json and args.output_json.resolve() == args.authoring_json.resolve():
        parser.error('Use separate files for the full context and authoring payloads.')
    authoring = {key: result[key] for key in ('headers', 'existingRows', 'addedRows', 'allRows', 'summary')}
    for requested, payload in ((args.output_json, result), (args.authoring_json, authoring)):
        if not requested:
            continue
        output = requested.resolve()
        protected = Path(result['summary']['reference']).resolve()
        if output == protected or output.suffix.casefold() != '.json':
            parser.error('--output-json must be a separate .json file, never the reference CSV.')
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_name('.' + output.name + '.pending')
        try:
            temporary.write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
            temporary.replace(output)
        finally:
            temporary.unlink(missing_ok=True)
    print(json.dumps(result['summary'], ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
